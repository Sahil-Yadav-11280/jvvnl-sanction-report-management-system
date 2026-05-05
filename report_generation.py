from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, Cm
from docx2pdf import convert
import os


def data_setting(document, key, value):
    data = document.add_paragraph()
    run1 = data.add_run(f'{key}: ')
    run1.bold = True
    run2 = data.add_run(f'{value}')

    for run in data.runs:
        run.font.size = Pt(12)
        run.font.name = 'Georgia'


def generate_report(
    report_id,
    book_no,
    page_no,
    purpose,
    date,
    customer_name,
    address,
    sanction_number,
    sanction_date,
    consumer_number,
    to_jen,
    proposed_work,
    ddr_number,
    ddr_date
):
    document = Document()

    para = document.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = para.add_run(f"Book No: {book_no}, Page No: {page_no}")
    run.font.size = Pt(12)
    run.font.name = 'Georgia'

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("Jaipur Vidhyut Vitaran Nigam Limited")
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = 'Georgia'

    order = document.add_paragraph()
    order.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = order.add_run("Sundry Job Order")
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = 'Georgia'

    pur = document.add_paragraph()
    pur.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = pur.add_run("Purpose: ")
    run1.bold = True
    run2 = pur.add_run(purpose)

    for run in pur.runs:
        run.font.size = Pt(12)
        run.font.name = 'Georgia'

    document.add_paragraph(" ")

    data_setting(document, 'Date', date)
    data_setting(document, 'Customer Name', customer_name)
    data_setting(document, 'Address', address)
    data_setting(document, 'Sanction Number', sanction_number)
    data_setting(document, 'Sanction Date', sanction_date)
    data_setting(document, 'Consumer Number', consumer_number)
    data_setting(document, 'To (JEN)', to_jen)

    document.add_paragraph(" ")

    table = document.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    title_cell = table.cell(0, 0)
    para = title_cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run('Proposed Work')
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = 'Georgia'

    data_cell = table.cell(1, 0)
    para = data_cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(proposed_work)
    run.font.size = Pt(12)
    run.font.name = 'Georgia'

    document.add_paragraph(" ")

    data_setting(document, 'Demand Deposit Receipt Number', ddr_number)
    data_setting(document, 'Demand Deposit Receipt Date', ddr_date)

    document.add_paragraph(" ")

    if os.path.exists("sig.jpg"):
        document.add_picture("sig.jpg", width=Cm(3))

    sig = document.add_paragraph()
    run = sig.add_run("Signature")
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = 'Georgia'

    docx_file = f"report_{report_id}.docx"
    pdf_file = f"report_{report_id}.pdf"

    document.save(docx_file)
    convert(docx_file, pdf_file)

    if os.path.exists(pdf_file):
        os.remove(docx_file)

    return pdf_file