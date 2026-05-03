from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt
from docx.shared import Cm
import os
from docx2pdf import convert

document = Document()

def data_setting(key , value):
    data = document.add_paragraph()
    run1 = data.add_run(f'{key}: ').bold = True
    run2 = data.add_run(f'{value}')
    for run in data.runs:
        run.font.size = Pt(12)
        run.font.name = 'georgia'


def generate_report(book_no , page_no , purpose , date , customer_name , address , sanction_number , sanction_date , consumer_number , to_jen , proposed_work ,ddr_number , ddr_date):
    # Book Number:
    book_no = document.add_paragraph()
    book_no.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_book = book_no.add_run(f"Book No: {book_no}, Page No: {page_no}")
    run_book.font.size = Pt(12)
    run_book.font.name = 'Georgia'

    # Heading:
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("Jaipur Vidhyut Vitaran Nigam Limited")
    run.font.bold = True
    run.font.size = Pt(18)
    run.font.name = 'Georgia'

    # Job Order:
    order = document.add_paragraph()
    order.alignment = WD_ALIGN_PARAGRAPH.CENTER
    order_run = order.add_run("Sundry Job Order")
    order_run.font.bold = True
    order_run.font.name = 'Georgia'
    order_run.font.size = Pt(12)

    #Purpose:
    pur = document.add_paragraph()
    pur.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pur_run = pur.add_run('Purpose: ').bold = True
    pur_run2 = pur.add_run(f'{purpose}')

    for run in pur.runs:
        run.font.size = Pt(12)
        run.font.name = 'Georgia'

    document.add_paragraph(" ")

    data_setting('Date' , f'{date}')

    data_setting('Customer Name' , f'{customer_name}')
    data_setting('Address' , f'{address}')
    data_setting('Sanction Number' , f'{sanction_number}')
    data_setting('Sanction Date' , f'{sanction_date}')
    data_setting('Consumer Number' , f'{consumer_number}')
    data_setting('To (JEN)' , f'{to_jen}')

    document.add_paragraph(" ")

    table = document.add_table(rows=2 , cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    table_title = table.cell(0,0)
    para = table_title.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run('Proposed Work')
    run.bold = True
    run.font.name = 'Georgia'
    run.font.size = Pt(12)

    table_data = table.cell(1,0)
    para = table_data.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(f'{proposed_work}')
    run.bold = False
    run.font.name = 'Georgia'
    run.font.size = Pt(12)

    document.add_paragraph(" ")

    data_setting('Demand Deposit Receipt Number' , f'{ddr_number}')
    data_setting('Demand Deposit Receipt Date' , f'{ddr_date}')

    document.add_paragraph(" ")

    document.add_picture('sig.jpg' , width=Cm(3))
    sig = document.add_paragraph()
    run = sig.add_run('Signature')
    run.font.size = Pt(12)
    run.font.name = 'Georgia'
    run.font.bold = True

    document.save('Test.docx')
    convert('Test.docx' , 'Test.pdf')
    if os.path.exists('Test.pdf'):
        os.remove('Test.docx')